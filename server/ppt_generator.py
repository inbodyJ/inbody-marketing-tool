from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ── InBody 공식 브랜드 컬러 ──────────────────
RED    = RGBColor(0x97, 0x1B, 0x2F)   # InBody Red
DGRAY  = RGBColor(0x4B, 0x4F, 0x5A)   # 짙은 회색 (본문)
CGRAY  = RGBColor(0x67, 0x76, 0x7F)   # Cool Gray (소제목)
IGRAY  = RGBColor(0xA2, 0xB2, 0xC8)   # Ice Gray (폴리곤)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BLACK  = RGBColor(0x10, 0x18, 0x20)

W = Inches(13.33)
H = Inches(7.5)


def _add_textbox(slide, l, t, w, h, text, size, bold=False,
                 color=DGRAY, align=PP_ALIGN.LEFT, italic=False):
    tx = slide.shapes.add_textbox(l, t, w, h)
    tf = tx.text_frame
    tf.word_wrap = True
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = "Arial"
    return tx


def _add_rect(slide, l, t, w, h, fill, line=False, line_color=IGRAY):
    from pptx.util import Pt as Pt2
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line_color
        shape.line.width = Pt2(0.5)
    else:
        shape.line.fill.background()
    return shape


def _add_iceberg_bg(slide):
    """우하단 Ice Gray 폴리곤 배경 — 공식 양식의 핵심 그래픽"""
    polys = [
        (Inches(10.8), Inches(5.2), Inches(2.5), Inches(2.3), 0.18),
        (Inches(11.6), Inches(4.6), Inches(1.7), Inches(2.9), 0.12),
        (Inches(12.2), Inches(5.8), Inches(1.1), Inches(1.7), 0.09),
    ]
    for l, t, w, h, opacity in polys:
        shape = slide.shapes.add_shape(1, l, t, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = IGRAY
        shape.fill.fore_color.theme_color  # 접근만 해도 됨
        shape.fill.fore_color.brightness = 0
        from pptx.oxml.ns import qn
        from lxml import etree
        solidFill = shape.fill._xPr.find(qn('a:solidFill'))
        if solidFill is not None:
            srgbClr = solidFill.find(qn('a:srgbClr'))
            if srgbClr is not None:
                alpha = etree.SubElement(srgbClr, qn('a:alpha'))
                alpha.set('val', str(int(opacity * 100000)))
        shape.line.fill.background()

    small_polys = [
        (Inches(0), Inches(0), Inches(1.2), Inches(1.1), 0.10),
        (Inches(0), Inches(0.5), Inches(0.7), Inches(0.8), 0.07),
    ]
    for l, t, w, h, opacity in small_polys:
        shape = slide.shapes.add_shape(1, l, t, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = IGRAY
        from pptx.oxml.ns import qn
        from lxml import etree
        solidFill = shape.fill._xPr.find(qn('a:solidFill'))
        if solidFill is not None:
            srgbClr = solidFill.find(qn('a:srgbClr'))
            if srgbClr is not None:
                alpha = etree.SubElement(srgbClr, qn('a:alpha'))
                alpha.set('val', str(int(opacity * 100000)))
        shape.line.fill.background()


def _add_footer(slide, page_num, total, ppt_title="InBody 제안서"):
    """공식 양식 하단 푸터: InBody 로고 + 중앙 선 + 페이지 번호"""
    _add_textbox(slide, Inches(0.4), Inches(7.0), Inches(1.5), Inches(0.4),
                 "InBody", 11, bold=True, color=DGRAY)
    line = slide.shapes.add_shape(1,
        Inches(6.2), Inches(7.15), Inches(0.9), Emu(12000))
    line.fill.solid()
    line.fill.fore_color.rgb = DGRAY
    line.line.fill.background()
    _add_textbox(slide, Inches(10.5), Inches(7.0), Inches(2.6), Inches(0.4),
                 f"{ppt_title}          {page_num}/{total}",
                 9, color=CGRAY, align=PP_ALIGN.RIGHT)


def _add_inbody_logo(slide):
    """우상단 InBody 로고 (빨간색)"""
    _add_textbox(slide, Inches(11.8), Inches(0.2), Inches(1.3), Inches(0.5),
                 "InBody", 14, bold=True, color=RED, align=PP_ALIGN.RIGHT)


def _add_slide_header(slide, number, title, subtitle=""):
    """본문 슬라이드 헤더: 01(빨간) + 대제목(짙은회색) + 소제목(빨간)"""
    tx = slide.shapes.add_textbox(Inches(0.4), Inches(0.28), Inches(10), Inches(0.65))
    tf = tx.text_frame
    tf.text = ""
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = f"{number:02d}  "
    r1.font.size = Pt(28)
    r1.font.bold = True
    r1.font.color.rgb = RED
    r1.font.name = "Arial"
    r2 = p.add_run()
    r2.text = title
    r2.font.size = Pt(28)
    r2.font.bold = True
    r2.font.color.rgb = DGRAY
    r2.font.name = "Arial"
    if subtitle:
        _add_textbox(slide, Inches(0.55), Inches(0.95), Inches(10), Inches(0.4),
                     subtitle, 14, color=CGRAY)


# ── 슬라이드 타입별 생성 함수 ────────────────

def slide_cover(prs, title, subtitle, author="", dept=""):
    """표지 슬라이드 — 슬라이드 1번 양식"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_iceberg_bg(slide)
    _add_inbody_logo(slide)
    _add_textbox(slide, Inches(0.6), Inches(2.8), Inches(9), Inches(1.2),
                 title, 40, bold=True, color=BLACK)
    if subtitle:
        _add_textbox(slide, Inches(0.6), Inches(4.05), Inches(9), Inches(0.5),
                     subtitle, 16, color=CGRAY)
    if author:
        _add_textbox(slide, Inches(0.6), Inches(6.2), Inches(4), Inches(0.35),
                     author, 14, bold=False, color=DGRAY)
    if dept:
        _add_textbox(slide, Inches(0.6), Inches(6.55), Inches(4), Inches(0.35),
                     dept, 12, color=CGRAY)


def slide_section(prs, part_num, title, subtitle=""):
    """섹션 구분 슬라이드 — Part 1 형식"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_iceberg_bg(slide)
    _add_textbox(slide, Inches(0), Inches(2.6), W, Inches(0.55),
                 f"Part {part_num}", 16, bold=True, color=RED,
                 align=PP_ALIGN.CENTER)
    _add_textbox(slide, Inches(0), Inches(3.1), W, Inches(1.0),
                 title, 44, bold=True, color=DGRAY,
                 align=PP_ALIGN.CENTER)
    if subtitle:
        _add_textbox(slide, Inches(0), Inches(4.15), W, Inches(0.5),
                     subtitle, 16, color=CGRAY, align=PP_ALIGN.CENTER)


def slide_text(prs, number, title, subtitle, body_text,
               page_num, total, ppt_title):
    """텍스트 본문 슬라이드"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_iceberg_bg(slide)
    _add_inbody_logo(slide)
    _add_slide_header(slide, number, title, subtitle)
    _add_footer(slide, page_num, total, ppt_title)

    if body_text:
        tx = slide.shapes.add_textbox(
            Inches(0.55), Inches(1.45), Inches(12.3), Inches(5.3))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.text = ""
        for line in (body_text if isinstance(body_text, list) else [body_text]):
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(6)
            run = p.add_run()
            run.text = str(line)
            run.font.size = Pt(16)
            run.font.color.rgb = DGRAY
            run.font.name = "Arial"


def slide_bullets(prs, number, title, subtitle, bullets,
                  page_num, total, ppt_title):
    """3단 카드 슬라이드 — 키워드 3개 배치"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_iceberg_bg(slide)
    _add_inbody_logo(slide)
    _add_slide_header(slide, number, title, subtitle)
    _add_footer(slide, page_num, total, ppt_title)

    items = bullets[:3] if isinstance(bullets, list) else [bullets]
    cols  = len(items)
    card_w = Inches(3.8) if cols == 3 else Inches(5.6)
    gap    = Inches(0.25)

    for i, item in enumerate(items):
        lft = Inches(0.55) + i * (card_w + gap)
        _add_rect(slide, lft, Inches(1.5), card_w, Inches(5.1),
                  WHITE, line=True, line_color=IGRAY)
        _add_rect(slide, lft, Inches(1.5), card_w, Emu(40000), RED)
        label = item if isinstance(item, str) else str(item)
        _add_textbox(slide, lft + Inches(0.2), Inches(1.65),
                     card_w - Inches(0.4), Inches(4.8),
                     label, 14, color=DGRAY)


def slide_thank_you(prs, contact="inbody.com"):
    """마지막 슬라이드 — Thank you"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_iceberg_bg(slide)
    _add_inbody_logo(slide)
    _add_textbox(slide, Inches(0), Inches(3.1), W, Inches(0.9),
                 "Thank you.", 44, bold=True, color=DGRAY,
                 align=PP_ALIGN.CENTER)
    if contact:
        _add_textbox(slide, Inches(0), Inches(4.1), W, Inches(0.5),
                     contact, 16, color=CGRAY, align=PP_ALIGN.CENTER)


# ── 메인 진입점 ─────────────────────────────
def create_ppt(slides_data: list, product: str) -> str:
    """
    slides_data 형식:
    [
      {"type": "cover",   "title": "...", "subtitle": "...", "author": "...", "dept": "..."},
      {"type": "section", "part": 1, "title": "...", "subtitle": "..."},
      {"type": "text",    "number": 1, "title": "...", "subtitle": "...", "body": ["..."]},
      {"type": "bullets", "number": 2, "title": "...", "subtitle": "...", "body": ["A","B","C"]},
      {"type": "thank",   "contact": "inbody.com"},
    ]
    """
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H

    total     = len(slides_data)
    ppt_title = product
    page_num  = 1

    for slide in slides_data:
        stype = slide.get("type", "text")

        if stype == "cover":
            slide_cover(prs,
                title=slide.get("title", product),
                subtitle=slide.get("subtitle", ""),
                author=slide.get("author", ""),
                dept=slide.get("dept", ""))

        elif stype == "section":
            slide_section(prs,
                part_num=slide.get("part", 1),
                title=slide.get("title", ""),
                subtitle=slide.get("subtitle", ""))

        elif stype == "bullets":
            slide_bullets(prs,
                number=slide.get("number", page_num),
                title=slide.get("title", ""),
                subtitle=slide.get("subtitle", ""),
                bullets=slide.get("body", []),
                page_num=page_num, total=total, ppt_title=ppt_title)
            page_num += 1

        elif stype == "thank":
            slide_thank_you(prs, contact=slide.get("contact", "inbody.com"))

        else:  # text
            slide_text(prs,
                number=slide.get("number", page_num),
                title=slide.get("title", ""),
                subtitle=slide.get("subtitle", ""),
                body_text=slide.get("body", []),
                page_num=page_num, total=total, ppt_title=ppt_title)
            page_num += 1

    output_dir = os.environ.get("TMPDIR", os.environ.get("TEMP", os.environ.get("TMP", "/tmp")))
    safe_name  = product.replace(" ", "_")
    filepath   = os.path.join(output_dir, f"InBody_{safe_name}_proposal.pptx")
    prs.save(filepath)
    return filepath
